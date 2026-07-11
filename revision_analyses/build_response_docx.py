#!/usr/bin/env python
"""Convert RESPONSE_TO_REVIEWERS.md into a formatted Word .docx (lightweight md->docx)."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path("/Users/danielliang/Library/CloudStorage/Dropbox/ClinVar_GenomMedicine/revision_analyses")
MD = BASE / "RESPONSE_TO_REVIEWERS.md"
OUT = BASE / "RESPONSE_TO_REVIEWERS.docx"

def add_runs(par, text):
    """Render inline **bold** and *italic* into a paragraph."""
    # split on bold/italic tokens
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            par.add_run(chunk[1:-1]).italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = par.add_run(chunk[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        else:
            par.add_run(chunk)

def main():
    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(10.5)
    lines = MD.read_text().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue
        # table block
        if ln.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]  # drop --- separator
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    for ci, cell in enumerate(r):
                        if ci < len(t.rows[ri].cells):
                            cpar = t.rows[ri].cells[ci].paragraphs[0]
                            add_runs(cpar, cell)
                            if ri == 0:
                                for rr in cpar.runs: rr.bold = True
            continue
        # headings
        if ln.startswith("# "):
            h = doc.add_heading(ln[2:], level=0)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:], level=2)
        elif ln.startswith("---"):
            doc.add_paragraph().add_run("_" * 60).font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        elif ln.lstrip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, ln.lstrip()[2:])
        elif ln.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            r = p.add_run(re.sub(r"^>\s?", "", ln)); r.italic = True; r.font.color.rgb = RGBColor(0x66,0x66,0x66)
        else:
            p = doc.add_paragraph(); add_runs(p, ln)
        i += 1
    doc.save(OUT)
    print(f"wrote {OUT}  ({OUT.stat().st_size:,} bytes)")

if __name__ == "__main__":
    main()
