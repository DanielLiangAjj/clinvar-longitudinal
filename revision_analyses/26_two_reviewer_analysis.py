# -*- coding: utf-8 -*-
"""
Two-reviewer analysis of the mechanism-validation sample (Reviewer 2, Major point #3).

Both co-authors adjudicated all 48 events independently, so this script reports what a single
workbook cannot: inter-rater reliability, and how much the heuristic's measured accuracy depends on
which reviewer you ask.

    python 26_two_reviewer_analysis.py

Estimators match 21_score_mechanism_validation.py exactly (same stratum weights, same NBOOT, same
SEED), so per-reviewer numbers here are directly comparable to that script's output.

Reads answers keyed by review_id rather than row position, so a workbook whose banner row was added
or deleted is handled correctly instead of silently shifting by one.

Outputs
-------
  outputs/two_reviewer_validation.json          machine-readable
  outputs/two_reviewer_validation.md            paste-ready summary
  Mechanism_Validation_Adjudication_Worklist.xlsx   the disagreements, for the two reviewers to settle
"""
import json
import math
import random
import collections
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REV = Path("/Users/danielliang/Library/CloudStorage/Dropbox/ClinVar_GenomMedicine/revision_analyses")
OUT = REV / "outputs"
WB_A = REV / "Mechanism_Validation_Review_FINAL_GZ_returned.xlsx"
WB_B = REV / "Mechanism_Validation_Review_FINAL_YL_returned.xlsx"
NAME_A, NAME_B = "Reviewer 1 (G.Z.)", "Reviewer 2 (Y.L.)"

DRIVERS = Path("/Users/danielliang/Downloads/ClinVar Research Project/data/reclassification_drivers.json")
if not DRIVERS.exists():
    DRIVERS = REV / "external_dependencies" / "ClinVar Research Project" / "data" / "reclassification_drivers.json"

LABELS = ["reinterpretation", "new_submitter", "conflict_emerged", "conflict_resolved",
          "batch_event", "likely_schema_driven", "likely_knowledge_driven", "transition_ambiguous"]
NBOOT = 20000
SEED = 20260720

ds = json.load(open(DRIVERS))["driver_summary"]
pop = {l: ds[l]["count"] for l in LABELS}
N = sum(pop.values())
w = {l: pop[l] / N for l in LABELS}


# ── read a workbook by review_id, immune to banner-row shifts ─────────────────
def read(path):
    wb = load_workbook(path)
    wf, wh = wb["Review form"], wb["Heuristic labels"]

    def scan(ws, cols):
        d = {}
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            if isinstance(k, str) and len(k) == 5 and k[0] == "R" and k[1:].isdigit():
                d[k] = {c: ws.cell(r, c).value for c in cols}
        return d

    form, heur = scan(wf, [2, 5, 6, 13, 14]), scan(wh, [2])
    if set(form) != set(heur):
        raise SystemExit(f"{path.name}: review_id sets differ between tabs")
    canon = {l.lower(): l for l in LABELS}
    out = {}
    for rid in sorted(form):
        lab = str(form[rid][13] or "").strip().lower()
        if lab and lab not in canon:
            raise SystemExit(f"{path.name} {rid}: '{lab}' is not one of the eight labels")
        out[rid] = {
            "gene": form[rid][2],
            "prev": str(form[rid][5] or ""),
            "curr": str(form[rid][6] or ""),
            "heuristic": str(heur[rid][2] or "").strip(),
            "label": canon.get(lab, ""),
            "confidence": str(form[rid][14] or "").strip(),
        }
    return out


# ── definitional admissibility ────────────────────────────────────────────────
# Two of the eight labels are only available for certain snapshots, and the workbook's
# 'Label definitions' tab says so explicitly. A label outside its admissible set is a
# mis-application of the taxonomy, not a judgment the heuristic can be wrong about.
BATCH_SNAPSHOTS = {"2013-12", "2014-01", "2014-02", "2018-06", "2020-01"}
TRANSITION = {"likely_schema_driven", "likely_knowledge_driven", "transition_ambiguous"}


def inadmissible(rec):
    """Return a reason string if this label cannot apply to this event, else None."""
    if rec["label"] == "batch_event" and rec["curr"] not in BATCH_SNAPSHOTS:
        return (f"batch_event requires a mass-update snapshot "
                f"({', '.join(sorted(BATCH_SNAPSHOTS))}); this event is {rec['curr']}")
    if rec["label"] in TRANSITION and not ("2020-01" <= rec["curr"] <= "2020-06"):
        return (f"{rec['label']} is defined only for the Jan–Jun 2020 transition window; "
                f"this event is {rec['curr']}")
    return None


A, B = read(WB_A), read(WB_B)
if set(A) != set(B):
    raise SystemExit("the two workbooks cover different events")
IDS = sorted(A)
for i in IDS:
    if A[i]["heuristic"] != B[i]["heuristic"]:
        raise SystemExit(f"{i}: heuristic label differs between workbooks — one has been altered")


# ── the 21_ estimator, verbatim in behaviour ──────────────────────────────────
def estimates(counts, n_c, active, wn):
    prec = {c: counts[c][c] / n_c[c] for c in active}
    acc = sum(wn[c] * prec[c] for c in active)
    prop = {t: sum(wn[c] * counts[c][t] / n_c[c] for c in active) for t in LABELS}
    rec = {}
    for t in LABELS:
        joint = wn[t] * counts[t][t] / n_c[t] if t in active else 0.0
        rec[t] = joint / prop[t] if prop[t] > 0 else float("nan")
    return acc, prec, prop, rec


def score(rows):
    """rows: list of (heuristic, adjudicated_label). Returns point estimates + bootstrap CIs."""
    cell = collections.defaultdict(collections.Counter)
    n_c = collections.Counter()
    for h, t in rows:
        cell[h][t] += 1
        n_c[h] += 1
    active = [l for l in LABELS if n_c[l] > 0]
    wsum = sum(w[l] for l in active)
    wn = {l: w[l] / wsum for l in active}
    acc, prec, prop, rec = estimates(cell, n_c, active, wn)

    strata = {c: [t for h, t in rows if h == c] for c in active}
    random.seed(SEED)
    bs = {"acc": [], "prec": collections.defaultdict(list), "prop": collections.defaultdict(list),
          "rec": collections.defaultdict(list)}
    for _ in range(NBOOT):
        cnt = collections.defaultdict(collections.Counter)
        nn = collections.Counter()
        for c in active:
            for t in random.choices(strata[c], k=n_c[c]):
                cnt[c][t] += 1
                nn[c] += 1
        a, p, q, rc = estimates(cnt, nn, active, wn)
        bs["acc"].append(a)
        for c in active:
            bs["prec"][c].append(p[c])
        for t in LABELS:
            bs["prop"][t].append(q[t])
            if rc[t] == rc[t]:
                bs["rec"][t].append(rc[t])
    k = sum(1 for h, t in rows if h == t)
    return {
        "n": len(rows), "k": k, "n_c": dict(n_c),
        "unweighted": (k / len(rows), *wilson(k, len(rows))[1:]),
        "weighted": (acc, *ci(bs["acc"])),
        "precision": {l: (prec.get(l), *ci(bs["prec"][l])) for l in LABELS},
        "recall": {l: (rec.get(l), *ci(bs["rec"][l])) for l in LABELS},
        "corrected": {l: (prop.get(l), *ci(bs["prop"][l])) for l in LABELS},
        "confusion": {c: dict(cell[c]) for c in LABELS},
    }


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"),) * 3
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return p, max(0.0, c - h), min(1.0, c + h)


def ci(v):
    if not v:
        return (float("nan"), float("nan"))
    s = sorted(v)
    return s[int(0.025 * len(s))], s[min(len(s) - 1, int(0.975 * len(s)))]


def pc(x):
    return "—" if x is None or x != x else f"{100 * x:.1f}"


# ── Cohen's kappa between the two reviewers ───────────────────────────────────
def kappa(pairs):
    n = len(pairs)
    po = sum(1 for a, b in pairs if a == b) / n
    ca = collections.Counter(a for a, _ in pairs)
    cb = collections.Counter(b for _, b in pairs)
    pe = sum((ca[l] / n) * (cb[l] / n) for l in LABELS)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan"), po, pe


pairs = [(A[i]["label"], B[i]["label"]) for i in IDS]
k_hat, po, pe = kappa(pairs)

# bootstrap kappa respecting the stratified design: resample within each heuristic stratum
by_stratum = collections.defaultdict(list)
for i in IDS:
    by_stratum[A[i]["heuristic"]].append((A[i]["label"], B[i]["label"]))
random.seed(SEED)
k_bs = []
for _ in range(NBOOT):
    samp = []
    for c, v in by_stratum.items():
        samp += random.choices(v, k=len(v))
    kk, _, _ = kappa(samp)
    if kk == kk:
        k_bs.append(kk)
k_lo, k_hi = ci(k_bs)

# ── per-reviewer and consensus scoring ────────────────────────────────────────
res_a = score([(A[i]["heuristic"], A[i]["label"]) for i in IDS])
res_b = score([(B[i]["heuristic"], B[i]["label"]) for i in IDS])

agreed = [i for i in IDS if A[i]["label"] == B[i]["label"]]
split = [i for i in IDS if A[i]["label"] != B[i]["label"]]
res_c = score([(A[i]["heuristic"], A[i]["label"]) for i in agreed])

# bounds on the fully adjudicated result: the split events could resolve either way
best = [(A[i]["heuristic"], A[i]["label"]) for i in agreed]
worst = list(best)
for i in split:
    h = A[i]["heuristic"]
    favour = h if h in (A[i]["label"], B[i]["label"]) else A[i]["label"]
    against = A[i]["label"] if A[i]["label"] != h else B[i]["label"]
    best.append((h, favour))
    worst.append((h, against))
res_best, res_worst = score(best), score(worst)

# ── sensitivity: drop events whose adjudicated label is definitionally inadmissible ────
bad_a = {i: inadmissible(A[i]) for i in IDS if inadmissible(A[i])}
bad_b = {i: inadmissible(B[i]) for i in IDS if inadmissible(B[i])}
keep_a = [i for i in IDS if i not in bad_a]
keep_b = [i for i in IDS if i not in bad_b]
res_a_adm = score([(A[i]["heuristic"], A[i]["label"]) for i in keep_a]) if len(keep_a) < len(IDS) else res_a
res_b_adm = score([(B[i]["heuristic"], B[i]["label"]) for i in keep_b]) if len(keep_b) < len(IDS) else res_b
both_adm = [i for i in IDS if i not in bad_a and i not in bad_b]
k_adm, po_adm, pe_adm = kappa([(A[i]["label"], B[i]["label"]) for i in both_adm])

print("=" * 82)
print("INTER-RATER RELIABILITY  (two blinded reviewers, all 48 events, 8 nominal labels)")
print("=" * 82)
print(f"Observed agreement      : {len(agreed)}/48 = {100*po:.1f}%")
print(f"Chance-expected         : {100*pe:.1f}%")
print(f"Cohen's kappa           : {k_hat:.3f} (bootstrap 95% CI {k_lo:.3f}–{k_hi:.3f})")

print("\n" + "=" * 82)
print("HEURISTIC ACCURACY BY REFERENCE STANDARD")
print("=" * 82)
print(f"{'reference standard':42s} {'n':>3s} {'unweighted':>18s} {'pop-weighted':>22s}")
for lbl, r in ((NAME_A, res_a), (NAME_B, res_b),
               (f"Both reviewers concur (n={len(agreed)})", res_c),
               ("Adjudicated — upper bound", res_best),
               ("Adjudicated — lower bound", res_worst)):
    u, ul, uh = r["unweighted"]
    g, gl, gh = r["weighted"]
    print(f"{lbl:42s} {r['n']:>3d} {r['k']}/{r['n']} {pc(u):>5s}% "
          f"[{pc(ul)}–{pc(uh)}] {pc(g):>7s}% [{pc(gl)}–{pc(gh)}]")

print("\n" + "=" * 82)
print("WHERE THE REVIEWERS SPLIT")
print("=" * 82)
print(f"{'id':6s} {'gene':10s} {'heuristic':24s} {NAME_A:24s} {NAME_B:24s}")
for i in split:
    tag = []
    if A[i]["label"] == A[i]["heuristic"]:
        tag.append("R1=heur")
    if B[i]["label"] == B[i]["heuristic"]:
        tag.append("R2=heur")
    print(f"{i:6s} {str(A[i]['gene']):10s} {A[i]['heuristic']:24s} "
          f"{A[i]['label']:24s} {B[i]['label']:24s} {'/'.join(tag)}")

print("\n" + "=" * 82)
print("DEFINITIONAL ADMISSIBILITY OF THE ADJUDICATED LABELS")
print("=" * 82)
for lbl, bad in ((NAME_A, bad_a), (NAME_B, bad_b)):
    print(f"{lbl}: {len(bad)} inadmissible label(s)")
    for i, why in bad.items():
        print(f"   {i} {str(A[i]['gene']):10s} assigned '{A[i]['label'] if bad is bad_a else B[i]['label']}'"
              f" — {why}")
        print(f"        heuristic said '{A[i]['heuristic']}'")
if bad_a or bad_b:
    print(f"\nExcluding those events (they cannot test the heuristic — the label was unavailable):")
    print(f"  {NAME_A:30s} {res_a_adm['k']}/{res_a_adm['n']} = {pc(res_a_adm['unweighted'][0])}% "
          f"unweighted, {pc(res_a_adm['weighted'][0])}% weighted "
          f"[{pc(res_a_adm['weighted'][1])}–{pc(res_a_adm['weighted'][2])}]")
    print(f"  {NAME_B:30s} {res_b_adm['k']}/{res_b_adm['n']} = {pc(res_b_adm['unweighted'][0])}% "
          f"unweighted, {pc(res_b_adm['weighted'][0])}% weighted "
          f"[{pc(res_b_adm['weighted'][1])}–{pc(res_b_adm['weighted'][2])}]")
    print(f"  Cohen's kappa on the {len(both_adm)} admissible events: {k_adm:.3f} "
          f"(observed agreement {100*po_adm:.1f}%)")

print("\n" + "=" * 82)
print("PER-LABEL PRECISION — how the two reference standards differ")
print("=" * 82)
print(f"{'mechanism':26s} {'pop share':>10s} {NAME_A:>20s} {NAME_B:>20s}")
for l in LABELS:
    pa = res_a["precision"][l][0]
    pb = res_b["precision"][l][0]
    print(f"{l:26s} {100*w[l]:>9.2f}% {pc(pa):>19s}% {pc(pb):>19s}%")

# ── artefacts ─────────────────────────────────────────────────────────────────
OUT.mkdir(exist_ok=True)
payload = {
    "reviewers": {"a": {"name": NAME_A, "workbook": WB_A.name},
                  "b": {"name": NAME_B, "workbook": WB_B.name}},
    "n_events": len(IDS), "population_events": N, "n_bootstrap": NBOOT, "seed": SEED,
    "inter_rater": {"observed_agreement": po, "expected_agreement": pe,
                    "cohens_kappa": k_hat, "kappa_ci": [k_lo, k_hi],
                    "n_agreed": len(agreed), "n_split": len(split)},
    "by_reference_standard": {
        "reviewer_a": res_a, "reviewer_b": res_b, "both_concur": res_c,
        "adjudicated_upper_bound": res_best, "adjudicated_lower_bound": res_worst},
    "split_events": [{"review_id": i, "gene": A[i]["gene"], "heuristic": A[i]["heuristic"],
                      "reviewer_a": A[i]["label"], "reviewer_b": B[i]["label"],
                      "conf_a": A[i]["confidence"], "conf_b": B[i]["confidence"]} for i in split],
    "admissibility": {
        "reviewer_a_inadmissible": {i: {"label": A[i]["label"], "curr": A[i]["curr"], "reason": r}
                                    for i, r in bad_a.items()},
        "reviewer_b_inadmissible": {i: {"label": B[i]["label"], "curr": B[i]["curr"], "reason": r}
                                    for i, r in bad_b.items()},
        "reviewer_a_admissible_only": res_a_adm,
        "reviewer_b_admissible_only": res_b_adm,
        "kappa_admissible_only": {"kappa": k_adm, "observed_agreement": po_adm,
                                  "n": len(both_adm)},
    },
}
(OUT / "two_reviewer_validation.json").write_text(json.dumps(payload, indent=2, default=str))

md = ["# Two-reviewer validation of the mechanism-attribution heuristic", "",
      f"<!-- generated by 26_two_reviewer_analysis.py; {NAME_A} = {WB_A.name}, "
      f"{NAME_B} = {WB_B.name} -->", "",
      f"Two co-authors independently adjudicated all {len(IDS)} sampled events, each blinded to the "
      f"label the heuristic had assigned.", "",
      "## Inter-rater reliability", "",
      f"- Observed agreement: **{len(agreed)}/{len(IDS)} = {100*po:.1f}%**",
      f"- Chance-expected agreement: {100*pe:.1f}%",
      f"- **Cohen's kappa = {k_hat:.3f}** (bootstrap 95% CI {k_lo:.3f}–{k_hi:.3f})", "",
      "## Heuristic accuracy depends on the reference standard", "",
      "| Reference standard | n | Unweighted | Population-weighted (95% CI) |", "|---|---:|---:|---:|"]
for lbl, r in ((NAME_A, res_a), (NAME_B, res_b), (f"Both reviewers concur", res_c),
               ("Adjudicated, upper bound", res_best), ("Adjudicated, lower bound", res_worst)):
    u, _, _ = r["unweighted"]
    g, gl, gh = r["weighted"]
    md.append(f"| {lbl} | {r['n']} | {r['k']}/{r['n']} = {pc(u)}% | {pc(g)}% ({pc(gl)}–{pc(gh)}) |")
md += ["", "## Per-label precision", "",
       "| Mechanism | Population share | " + NAME_A + " | " + NAME_B + " |", "|---|---:|---:|---:|"]
for l in LABELS:
    md.append(f"| `{l}` | {100*w[l]:.2f}% | {pc(res_a['precision'][l][0])}% | "
              f"{pc(res_b['precision'][l][0])}% |")
if bad_a or bad_b:
    md += ["", "## Definitional admissibility", "",
           "`batch_event` is defined on the `Label definitions` tab as applying only to the mass-update "
           "snapshots 2013-12, 2014-01, 2014-02, 2018-06 and 2020-01; the three transition labels apply "
           "only to Jan–Jun 2020. A label outside its admissible set cannot test the heuristic, because "
           "that label was not available for that event.", ""]
    for lbl, bad, R in ((NAME_A, bad_a, A), (NAME_B, bad_b, B)):
        md.append(f"- **{lbl}: {len(bad)} inadmissible.**"
                  + ("" if not bad else " " + "; ".join(
                      f"`{i}` ({R[i]['gene']}, {R[i]['curr']}) assigned `{R[i]['label']}`" for i in bad)))
    md += ["", "| Excluding inadmissible events | n | Unweighted | Population-weighted (95% CI) |",
           "|---|---:|---:|---:|"]
    for lbl, r in ((NAME_A, res_a_adm), (NAME_B, res_b_adm)):
        u, _, _ = r["unweighted"]
        g, gl, gh = r["weighted"]
        md.append(f"| {lbl} | {r['n']} | {r['k']}/{r['n']} = {pc(u)}% | {pc(g)}% ({pc(gl)}–{pc(gh)}) |")
    md += ["", f"Cohen's kappa on the {len(both_adm)} events admissible for both reviewers: "
               f"**{k_adm:.3f}** (observed agreement {100*po_adm:.1f}%).", ""]

md += ["", f"## The {len(split)} events the reviewers judged differently", "",
       f"| Event | Gene | Heuristic | {NAME_A} | {NAME_B} |", "|---|---|---|---|---|"]
for i in split:
    md.append(f"| {i} | {A[i]['gene']} | `{A[i]['heuristic']}` | `{A[i]['label']}` | "
              f"`{B[i]['label']}` |")
md += ["", "These require adjudication by discussion before a single consensus figure can be quoted. "
           "Until then the defensible reporting is either both reviewers separately with the kappa, or "
           "the bounds above.", ""]
(OUT / "two_reviewer_validation.md").write_text("\n".join(md))

# adjudication worklist
wb = load_workbook(WB_A)
src = wb["Review form"]
hdr_row = next(r for r in range(1, src.max_row + 1)
               if str(src.cell(r, 1).value).strip() == "review_id")
keep = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
out = load_workbook(WB_A)
for s in out.sheetnames:
    if s not in ("Label definitions", "SCV evidence"):
        del out[s]
ws = out.create_sheet("Adjudicate", 0)
ws.cell(1, 1).value = (f"{len(split)} events where the two blinded reviewers disagreed. Settle each by "
                       "discussion and enter one agreed mechanism in the 'CONSENSUS' column. Do not "
                       "sort or delete rows. Evidence for every event is on the 'SCV evidence' tab.")
ws.cell(1, 1).font = Font(bold=True)
cols = ([str(src.cell(hdr_row, c).value) for c in keep]
        + [NAME_A, f"conf ({NAME_A})", NAME_B, f"conf ({NAME_B})", "Heuristic label", "CONSENSUS"])
for c, name in enumerate(cols, 1):
    cell = ws.cell(2, c)
    cell.value = name
    cell.font = Font(bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
yellow = PatternFill("solid", fgColor="FFF2CC")
srcmap = {}
for r in range(1, src.max_row + 1):
    k = src.cell(r, 1).value
    if isinstance(k, str) and len(k) == 5 and k[0] == "R":
        srcmap[k] = r
for n, i in enumerate(split, start=3):
    for c, sc in enumerate(keep, 1):
        ws.cell(n, c).value = src.cell(srcmap[i], sc).value
    base = len(keep)
    for off, v in enumerate([A[i]["label"], A[i]["confidence"], B[i]["label"], B[i]["confidence"],
                             A[i]["heuristic"]], 1):
        ws.cell(n, base + off).value = v
    ws.cell(n, base + 6).fill = yellow
for c in range(1, len(cols) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 15 if c > 4 else 12
out.save(REV / "Mechanism_Validation_Adjudication_Worklist.xlsx")

print(f"\nwrote outputs/two_reviewer_validation.json")
print(f"wrote outputs/two_reviewer_validation.md")
print(f"wrote Mechanism_Validation_Adjudication_Worklist.xlsx ({len(split)} events)")
