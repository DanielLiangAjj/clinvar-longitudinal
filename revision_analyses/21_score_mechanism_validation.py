# -*- coding: utf-8 -*-
"""
Score the completed mechanism-validation workbook (Reviewer 2, Major point #3).

Run this once co-authors return Mechanism_Validation_Review_FINAL.xlsx.

    python 21_score_mechanism_validation.py [path/to/filled_workbook.xlsx]

Why this supersedes 09_validation_score.py
------------------------------------------
The sample is STRATIFIED — six events drawn from each of the eight heuristic labels — not a simple
random sample of the 697,698 events. Raw correct/48 therefore estimates accuracy on a
label-balanced sample, which is not the population accuracy and would be wrong to report as such.
Because the strata are defined by the heuristic's own output, the design does support the right
estimands directly:

  precision_c   P(true = c | predicted = c)      read straight off stratum c
  P(true = t)   Σ_c w_c · (n_ct / n_c)           the corrected mechanism proportion
  recall_t      w_t·(n_tt/n_t) / P(true = t)
  accuracy      Σ_c w_c · precision_c            population-weighted

where w_c is the population share of label c. Confidence intervals come from a parametric
bootstrap that resamples each stratum multinomially, which behaves better than a normal
approximation at n_c = 6.

The corrected proportions with intervals are precisely the "misclassification rates or uncertainty
intervals for the estimated mechanism proportions" Reviewer 2 asked for.

Outputs  (named after the workbook, so two reviewers' runs cannot overwrite each other)
-------
  outputs/mechanism_validation_results__<workbook stem>.json          machine-readable
  outputs/Supplementary_Table_mechanism_validation__<workbook stem>.md   paste-ready table

For the two-reviewer analysis actually reported in the paper, see 26_two_reviewer_analysis.py.
This script scores ONE workbook against the heuristic and knows nothing about a second reviewer.
"""
import sys, json, math, collections, random, re
from pathlib import Path
from openpyxl import load_workbook

REV = Path("/Users/danielliang/Library/CloudStorage/Dropbox/ClinVar_GenomMedicine/revision_analyses")
PROJ = Path("/Users/danielliang/Downloads/ClinVar Research Project")
WB = Path(sys.argv[1]) if len(sys.argv) > 1 else REV / "Mechanism_Validation_Review_FINAL.xlsx"
OUT = REV / "outputs"
DRIVERS = PROJ / "data" / "reclassification_drivers.json"
# The parent project is only present on the Mac mini. revision_analyses/external_dependencies
# carries a verified snapshot of the one file needed here, so the scorer runs on any machine.
if not DRIVERS.exists():
    DRIVERS = REV / "external_dependencies" / "ClinVar Research Project" / "data" / "reclassification_drivers.json"
if not DRIVERS.exists():
    sys.exit(f"ERROR: cannot find reclassification_drivers.json (looked in {PROJ} and external_dependencies/)")

LABELS = ["reinterpretation", "new_submitter", "conflict_emerged", "conflict_resolved",
          "batch_event", "likely_schema_driven", "likely_knowledge_driven",
          "transition_ambiguous"]
NBOOT = 20000
SEED = 20260720


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"),) * 3
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return p, max(0.0, c - h), min(1.0, c + h)


def pct(x):
    return "—" if x != x else f"{100 * x:.1f}"


# ── read the workbook ─────────────────────────────────────────────────────────
wb = load_workbook(WB, data_only=False)
if "Review form" not in wb.sheetnames or "Heuristic labels" not in wb.sheetnames:
    sys.exit(f"ERROR: {WB.name} is not the validation workbook (tabs: {wb.sheetnames})")
wf, wh = wb["Review form"], wb["Heuristic labels"]

rows = []
r = 3
while wf.cell(r, 1).value:
    rid = str(wf.cell(r, 1).value).strip()
    hid = str(wh.cell(r, 1).value or "").strip()
    if rid != hid:
        sys.exit(f"ERROR: row {r} is misaligned — Review form has {rid}, Heuristic labels has {hid}. "
                 "Rows must not be sorted or inserted; re-download a clean workbook.")
    rows.append({
        "review_id": rid,
        "gene": wf.cell(r, 2).value,
        "heuristic": str(wh.cell(r, 2).value or "").strip(),
        # agreement is recomputed here, never read from the sheet's formula
        "truth": str(wf.cell(r, 13).value or "").strip(),
        "confidence": str(wf.cell(r, 14).value or "").strip(),
        "notes": str(wf.cell(r, 15).value or "").strip(),
        "row": r,
    })
    r += 1

# ── integrity checks before any statistics ────────────────────────────────────
problems = []
excluded = set()          # rows whose mechanism is unusable; these cannot be scored
for x in rows:
    if x["truth"] and x["truth"].lower() not in {l.lower() for l in LABELS}:
        problems.append(f"row {x['row']} ({x['review_id']}): mechanism '{x['truth']}' is not one of "
                        f"the eight labels — this event is EXCLUDED from the statistics")
        excluded.add(x["row"])
    # A bad confidence value is a nuisance, not a reason to discard a valid mechanism judgment;
    # it only removes the row from the high-confidence subgroup.
    if x["confidence"] and x["confidence"].lower() not in ("high", "medium", "low"):
        problems.append(f"row {x['row']} ({x['review_id']}): confidence '{x['confidence']}' is not "
                        f"High/Medium/Low — the mechanism still counts, the row is just left out "
                        f"of the high-confidence subgroup")

canon = {l.lower(): l for l in LABELS}
for x in rows:
    if x["truth"] and x["truth"].lower() in canon:
        x["truth"] = canon[x["truth"].lower()]

filled = [x for x in rows if x["truth"] and x["row"] not in excluded]
print(f"Workbook : {WB.name}")
print(f"Events   : {len(rows)} total, {len(filled)} adjudicated ({100*len(filled)/max(1,len(rows)):.0f}%)")
if problems:
    print("\n!! FIX THESE BEFORE REPORTING:")
    for p in problems:
        print("   -", p)
if not filled:
    print("\nNothing scored yet. Fill the 'YOUR ANSWER: mechanism' column and re-run.")
    sys.exit(0)
missing = [x["review_id"] for x in rows if not x["truth"] or x["row"] in excluded]
if missing:
    print(f"Not yet adjudicated ({len(missing)}): {', '.join(missing)}")
    print("Numbers below are provisional until every event is filled in.")

# ── population weights ────────────────────────────────────────────────────────
ds = json.load(open(DRIVERS))["driver_summary"]
pop = {l: ds.get(l, {}).get("count", 0) for l in LABELS}
N = sum(pop.values())
w = {l: pop[l] / N for l in LABELS}

# ── stratum tallies: n_ct = adjudicated events with heuristic c whose truth is t ──
cell = collections.defaultdict(collections.Counter)
n_c = collections.Counter()
for x in filled:
    cell[x["heuristic"]][x["truth"]] += 1
    n_c[x["heuristic"]] += 1

incomplete = [l for l in LABELS if n_c[l] == 0]
if incomplete:
    print(f"\nWARNING: no adjudicated events for label(s) {incomplete}; "
          "weighted estimates below omit those strata and are biased. Complete them.")

active = [l for l in LABELS if n_c[l] > 0]
wsum = sum(w[l] for l in active)
wn = {l: w[l] / wsum for l in active}          # renormalised over strata that have data


def estimates(counts):
    """Point estimates from a table counts[c][t]. Returns (accuracy, precision, proportion, recall)."""
    prec = {c: counts[c][c] / n_c[c] for c in active}
    acc = sum(wn[c] * prec[c] for c in active)
    prop = {t: sum(wn[c] * counts[c][t] / n_c[c] for c in active) for t in LABELS}
    rec = {}
    for t in LABELS:
        joint = wn[t] * counts[t][t] / n_c[t] if t in active else 0.0
        rec[t] = joint / prop[t] if prop[t] > 0 else float("nan")
    return acc, prec, prop, rec


acc_hat, prec_hat, prop_hat, rec_hat = estimates(cell)

# ── parametric bootstrap: resample each stratum multinomially ─────────────────
random.seed(SEED)
strata = {c: [x["truth"] for x in filled if x["heuristic"] == c] for c in active}
bs = {"acc": [], "prec": collections.defaultdict(list),
      "prop": collections.defaultdict(list), "rec": collections.defaultdict(list)}
for _ in range(NBOOT):
    cnt = collections.defaultdict(collections.Counter)
    for c in active:
        for t in random.choices(strata[c], k=n_c[c]):
            cnt[c][t] += 1
    a, p, q, rc = estimates(cnt)
    bs["acc"].append(a)
    for c in active:
        bs["prec"][c].append(p[c])
    for t in LABELS:
        bs["prop"][t].append(q[t])
        if rc[t] == rc[t]:
            bs["rec"][t].append(rc[t])


def ci(v):
    if not v:
        return (float("nan"), float("nan"))
    s = sorted(v)
    return s[int(0.025 * len(s))], s[min(len(s) - 1, int(0.975 * len(s)))]


acc_lo, acc_hi = ci(bs["acc"])
raw_correct = sum(1 for x in filled if x["truth"] == x["heuristic"])
raw_p, raw_lo, raw_hi = wilson(raw_correct, len(filled))

print("\n" + "=" * 78)
print("HEURISTIC ACCURACY")
print("=" * 78)
print(f"Unweighted, on the label-balanced sample : {raw_correct}/{len(filled)} = {pct(raw_p)}% "
      f"(Wilson 95% CI {pct(raw_lo)}–{pct(raw_hi)})")
print(f"Population-weighted (report this one)    : {pct(acc_hat)}% "
      f"(bootstrap 95% CI {pct(acc_lo)}–{pct(acc_hi)})")

hi_conf = [x for x in filled if x["confidence"].lower() == "high"]
if hi_conf:
    k = sum(1 for x in hi_conf if x["truth"] == x["heuristic"])
    p, lo, hi = wilson(k, len(hi_conf))
    print(f"Unweighted, high-confidence rows only    : {k}/{len(hi_conf)} = {pct(p)}% "
          f"(Wilson 95% CI {pct(lo)}–{pct(hi)})")

print("\n" + "=" * 78)
print("PER-LABEL PRECISION AND RECALL")
print("=" * 78)
print(f"{'mechanism':26s} {'n':>3s} {'precision':>22s} {'recall':>10s}")
for l in LABELS:
    if l not in active:
        print(f"{l:26s} {0:>3d} {'not adjudicated':>22s}")
        continue
    plo, phi = ci(bs["prec"][l])
    rlo, rhi = ci(bs["rec"][l])
    print(f"{l:26s} {n_c[l]:>3d} {pct(prec_hat[l]):>8s}% [{pct(plo)}–{pct(phi)}] "
          f"{pct(rec_hat[l]):>7s}% [{pct(rlo)}–{pct(rhi)}]")

print("\n" + "=" * 78)
print("MECHANISM PROPORTIONS — heuristic vs validation-corrected  (Reviewer 2's request)")
print("=" * 78)
print(f"{'mechanism':26s} {'heuristic':>10s} {'corrected':>10s} {'95% CI':>18s}")
table = []
for l in LABELS:
    lo, hi = ci(bs["prop"][l])
    print(f"{l:26s} {100*w[l]:>9.2f}% {pct(prop_hat[l]):>9s}% "
          f"{'['+pct(lo)+'–'+pct(hi)+']':>18s}")
    table.append({"mechanism": l, "population_n": pop[l], "heuristic_pct": 100 * w[l],
                  "corrected_pct": 100 * prop_hat[l], "corrected_lo": 100 * lo,
                  "corrected_hi": 100 * hi,
                  "precision": prec_hat.get(l), "recall": rec_hat.get(l),
                  "n_adjudicated": n_c[l]})

print("\n" + "=" * 78)
print("CONFUSION MATRIX  (rows = heuristic label, columns = reviewer's judgment)")
print("=" * 78)
hdr = "".join(f"{l[:11]:>12s}" for l in LABELS)
print(f"{'':26s}{hdr}")
for c in LABELS:
    line = "".join(f"{cell[c][t] if cell[c][t] else '·':>12}" for t in LABELS)
    print(f"{c:26s}{line}")

disagree = [x for x in filled if x["truth"] != x["heuristic"]]
if disagree:
    print(f"\nDISAGREEMENTS ({len(disagree)}) — the rows worth reading in the response letter")
    for x in disagree:
        print(f"  {x['review_id']} {str(x['gene']):10s} heuristic={x['heuristic']:24s} "
              f"reviewer={x['truth']:24s} conf={x['confidence'] or '—'}")
        if x["notes"]:
            print(f"       note: {x['notes'][:110]}")

# ── artefacts ─────────────────────────────────────────────────────────────────
OUT.mkdir(exist_ok=True)
res = {
    "workbook": WB.name, "n_events": len(rows), "n_adjudicated": len(filled),
    "population_events": N, "n_bootstrap": NBOOT, "seed": SEED,
    "accuracy_unweighted": {"k": raw_correct, "n": len(filled), "pct": 100 * raw_p,
                            "ci": [100 * raw_lo, 100 * raw_hi], "method": "Wilson"},
    "accuracy_weighted": {"pct": 100 * acc_hat, "ci": [100 * acc_lo, 100 * acc_hi],
                          "method": "stratum-weighted, parametric bootstrap"},
    "per_label": table,
    "confusion": {c: dict(cell[c]) for c in LABELS},
    "disagreements": [{"review_id": x["review_id"], "gene": x["gene"], "heuristic": x["heuristic"],
                       "reviewer": x["truth"], "confidence": x["confidence"], "notes": x["notes"]}
                      for x in disagree],
    "unadjudicated": missing, "integrity_problems": problems,
}
# Output names carry the workbook they came from. Two reviewers returned workbooks in this project,
# and a fixed filename let the second run silently overwrite the first — producing a results file
# whose numbers belonged to a different reviewer than the documents built from it. Never do that again.
slug = re.sub(r"[^A-Za-z0-9]+", "_", WB.stem).strip("_")
res["source_workbook_path"] = str(WB)
(OUT / f"mechanism_validation_results__{slug}.json").write_text(json.dumps(res, indent=2))

md = ["# Supplementary Table — validation of the reclassification-mechanism heuristic", "",
      f"<!-- generated by 21_score_mechanism_validation.py from {WB.name}; "
      f"{len(filled)}/{len(rows)} events adjudicated -->", ""]
if missing:
    md += [f"> **PROVISIONAL — {len(missing)} of {len(rows)} events are not yet adjudicated "
           f"({', '.join(missing)}). Do not cite these numbers until the review is complete.**", ""]
md += [f"Stratified random sample of {len(rows)} reclassification events (six per mechanism label, "
      f"fixed seed), adjudicated by manual review of the submission-level (SCV) record at the two "
      f"snapshots bracketing each event. Reviewers were blinded to the heuristic's label. "
      f"Estimates are weighted by each label's share of the {N:,} events; intervals are "
      f"{NBOOT:,}-draw parametric bootstrap percentiles.", "",
      "| Mechanism | Events (n) | Heuristic share | Precision (95% CI) | Recall (95% CI) | "
      "Corrected share (95% CI) |", "|---|---:|---:|---:|---:|---:|"]
for t in table:
    l = t["mechanism"]
    if n_c[l] == 0:
        md.append(f"| `{l}` | {t['population_n']:,} | {t['heuristic_pct']:.2f}% | — | — | — |")
        continue
    plo, phi = ci(bs["prec"][l]); rlo, rhi = ci(bs["rec"][l])
    md.append(f"| `{l}` | {t['population_n']:,} | {t['heuristic_pct']:.2f}% | "
              f"{pct(prec_hat[l])}% ({pct(plo)}–{pct(phi)}) | "
              f"{pct(rec_hat[l])}% ({pct(rlo)}–{pct(rhi)}) | "
              f"{t['corrected_pct']:.2f}% ({t['corrected_lo']:.2f}–{t['corrected_hi']:.2f}) |")
md += ["", f"**Overall agreement**, weighted to the population of events: "
           f"{pct(acc_hat)}% (95% CI {pct(acc_lo)}–{pct(acc_hi)}). "
           f"Unweighted agreement on the label-balanced sample was {raw_correct}/{len(filled)} "
           f"= {pct(raw_p)}% (Wilson 95% CI {pct(raw_lo)}–{pct(raw_hi)})."]
(OUT / f"Supplementary_Table_mechanism_validation__{slug}.md").write_text("\n".join(md))

print(f"\nwrote outputs/mechanism_validation_results__{slug}.json")
print(f"wrote outputs/Supplementary_Table_mechanism_validation__{slug}.md")
