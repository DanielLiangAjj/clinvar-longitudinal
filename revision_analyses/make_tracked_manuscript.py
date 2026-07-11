#!/usr/bin/env python
"""
Produce a Word manuscript with EDITS AS SUGGESTIONS (tracked changes) + rationale COMMENTS.

Works on a COPY (manuscript_tracked_changes.docx); never touches the original.
Each edit becomes a w:del (old text struck) + w:ins (new text) minimal-diff change, with a Word
comment carrying the rationale. New content (terminology box, tiered/stats/ITS paragraphs,
Acknowledgements) is inserted as tracked w:ins paragraphs.

Reviewer can Accept/Reject each suggestion in Word and read the rationale in the comment pane.
"""
import copy, sys
from docx import Document
from docx.text.run import Run
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = "/Users/danielliang/Library/CloudStorage/Dropbox/ClinVar_GenomMedicine/revision_analyses/manuscript_tracked_changes.docx"
AUTHOR = "Claude (ClinVar revision)"; INITIALS = "CR"; DATE = "2026-07-11T00:00:00Z"
_cid = [900]
def _nid():
    _cid[0] += 1; return _cid[0]

def _set_track_attrs(el):
    el.set(qn('w:id'), str(_nid())); el.set(qn('w:author'), AUTHOR); el.set(qn('w:date'), DATE)

def _new_wt(text):
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; return t

def split_run_at(para, offset):
    """Ensure a run boundary at char `offset`; split the containing run (formatting preserved)."""
    pos = 0
    for run in para.runs:
        rlen = len(run.text)
        if pos < offset < pos + rlen:
            left, right = run.text[:offset - pos], run.text[offset - pos:]
            newr = copy.deepcopy(run._r)
            for t in newr.findall(qn('w:t')):
                newr.remove(t)
            newr.append(_new_wt(right))
            run.text = left
            run._r.addnext(newr)
            return
        pos += rlen

def runs_with_offsets(para):
    out, pos = [], 0
    for run in para.runs:
        out.append((run, pos, pos + len(run.text))); pos += len(run.text)
    return out

def make_ins_run(text, template_run):
    """Create a run (copying template formatting) wrapped in w:ins. Returns Run object."""
    r = copy.deepcopy(template_run._r) if template_run is not None else OxmlElement('w:r')
    for t in r.findall(qn('w:t')):
        r.remove(t)
    for t in r.findall(qn('w:delText')):
        r.remove(t)
    r.append(_new_wt(text))
    ins = OxmlElement('w:ins'); _set_track_attrs(ins)
    ins.append(r)
    return ins, r

def wrap_del(run_el):
    for t in run_el.findall(qn('w:t')):
        t.tag = qn('w:delText')
    d = OxmlElement('w:del'); _set_track_attrs(d)
    run_el.addprevious(d); d.append(run_el)

def _runtext(para):
    return "".join(r.text for r in para.runs)

def tracked_replace(doc, find, replace, comment):
    """Minimal-diff tracked replacement of `find`->`replace`, operating in RUN-TEXT space so any
    non-run content (cross-reference fields, hyperlinks) in the paragraph is left untouched."""
    paras = [p for p in doc.paragraphs if find in _runtext(p)]
    if not paras:
        # fall back to para.text match -> whole-paragraph replacement (rare)
        alt = [p for p in doc.paragraphs if find in p.text]
        if alt:
            return whole_para_replace(doc, alt[0], alt[0].text.replace(find, replace, 1), comment)
        print(f"  [SKIP] not found: {find[:55]!r}"); return False
    para = paras[0]
    old = _runtext(para)
    new = old.replace(find, replace, 1)
    # common prefix / suffix -> localize change
    pl = 0
    while pl < len(old) and pl < len(new) and old[pl] == new[pl]:
        pl += 1
    sl = 0
    while sl < min(len(old), len(new)) - pl and old[-1 - sl] == new[-1 - sl]:
        sl += 1
    old_mid, new_mid = old[pl:len(old) - sl], new[pl:len(new) - sl]

    runs_ok = sum(len(r.text) for r in para.runs) == len(old)
    if not runs_ok:
        return whole_para_replace(doc, para, new, comment)

    split_run_at(para, pl)
    split_run_at(para, len(old) - sl)
    ro = runs_with_offsets(para)
    mid_runs = [r for (r, s, e) in ro if s >= pl and e <= len(old) - sl and s < e and s >= pl and e <= (len(old) - sl)]
    mid_runs = [r for (r, s, e) in ro if pl <= s and e <= (len(old) - sl) and (s < e)]
    template = mid_runs[0] if mid_runs else (para.runs[0] if para.runs else None)
    # anchor: first mid run (to insert before), else the run starting at pl
    anchor_el = mid_runs[0]._r if mid_runs else None
    if anchor_el is None:
        for (r, s, e) in ro:
            if s == pl:
                anchor_el = r._r; break
    ins_run_obj = None
    if new_mid:
        ins_el, ins_r = make_ins_run(new_mid, template)
        if anchor_el is not None:
            anchor_el.addprevious(ins_el)
        else:
            para._p.append(ins_el)
        ins_run_obj = Run(ins_r, para)
    for r in mid_runs:
        wrap_del(r._r)
    # comment on the inserted run (or first deleted run)
    tgt = ins_run_obj if ins_run_obj is not None else (Run(mid_runs[0]._r, para) if mid_runs else None)
    if tgt is not None and comment:
        doc.add_comment([tgt], text=comment, author=AUTHOR, initials=INITIALS)
    print(f"  [ok] replace: {find[:50]!r}")
    return True

def whole_para_replace(doc, para, new_text, comment):
    """Fallback: delete all inline content, insert new_text (tracked). Used when the paragraph
    has non-run content (fields/hyperlinks) so offset mapping is unreliable."""
    old = para.text
    p = para._p
    # delete every inline child (runs, hyperlinks) by wrapping runs; simplest: strike whole text
    for child in list(p):
        if child.tag in (qn('w:r'), qn('w:hyperlink')):
            p.remove(child)
    del_run = OxmlElement('w:r'); del_run.append(_new_wt(old)); del_run.find(qn('w:t')).tag = qn('w:delText')
    dele = OxmlElement('w:del'); _set_track_attrs(dele); dele.append(del_run)
    p.append(dele)
    ins_el, ins_r = make_ins_run(new_text, None)
    p.append(ins_el)
    if comment:
        doc.add_comment([Run(ins_r, para)], text=comment, author=AUTHOR, initials=INITIALS)
    print(f"  [ok] whole-para replace (had fields): {old[:45]!r}")
    return True

def insert_after(doc, anchor_substr, new_paragraphs, comment):
    """Insert new tracked-insertion paragraphs after the paragraph containing anchor_substr."""
    paras = [p for p in doc.paragraphs if anchor_substr in p.text]
    if not paras:
        print(f"  [SKIP insert] anchor not found: {anchor_substr[:45]!r}"); return False
    anchor = paras[0]
    prev_el = anchor._p
    first_run_obj = None
    for text in new_paragraphs:
        newp = OxmlElement('w:p')
        ins_el, ins_r = make_ins_run(text, None)
        # mark paragraph mark as inserted too (pPr/rPr/ins) so the whole new para is a suggestion
        ppr = OxmlElement('w:pPr'); rpr = OxmlElement('w:rPr'); pins = OxmlElement('w:ins'); _set_track_attrs(pins)
        rpr.append(pins); ppr.append(rpr); newp.append(ppr)
        newp.append(ins_el)
        prev_el.addnext(newp); prev_el = newp
        if first_run_obj is None:
            first_run_obj = Run(ins_r, Paragraph(newp, anchor._parent))
    if comment and first_run_obj is not None:
        doc.add_comment([first_run_obj], text=comment, author=AUTHOR, initials=INITIALS)
    print(f"  [ok] inserted {len(new_paragraphs)} para(s) after {anchor_substr[:35]!r}")
    return True

# ------------------------------------------------------------------ edits
REPLACEMENTS = [
 ("Overall, 9.3% of variants underwent at least one classification change, including 6,340 drastic pathogenic-to-benign (or vice versa) switches.",
  "Overall, 9.3% of variants underwent at least one change in aggregate classification category; however, only 0.68% (25,105 variants) ever underwent a medically significant change crossing the clinical actionability boundary, including 6,340 drastic pathogenic-to-benign (or vice versa) switches. Accounting for censoring, the cumulative incidence of first reclassification was 2.5% at one year and 15.5% at five years after a variant’s first appearance.",
  "R1 major & R2 major #1: the 9.3% counts ANY category change. We now distinguish the clinically significant subset (0.68%, 25,105 variants) and add the censored cumulative incidence (R2 #2). Verified directly: 0 of 697,698 events are category-stable, so concordant new submissions are NOT inflating this count."),
 ("Of 697,698 reclassification events, 33.2% resulted from submitter reinterpretation and 18.3% from new submitters.",
  "Of 697,698 reclassification events, 33.2% resulted from submitter reinterpretation and 18.3% from new-submitter-driven category shifts.",
  "R1 major: 'new submitter' events DO change the consensus category, so the earlier 'concordant new submissions' wording was misleading; relabeled."),
 ("A variant is considered reclassified if its categorization changed between any two consecutive snapshots, though we note that adjacent transitions (e.g., LP to P, LB to B) are clinically less significant than cross-tier transitions (e.g., LB to LP). This includes all aggregate classification changes, including adjacent-category transitions (e.g., LP→P) that, while not clinically distinct in actionability, reflect updated evidence strength.",
  "A variant is considered reclassified if its normalized consensus category changed between any two consecutive snapshots. This includes adjacent within-tier transitions (e.g., LP→P, LB→B) that ClinVar treats as concordant; we therefore stratify reclassifications by clinical significance in Section 3.2.1 rather than treating all category changes as equivalent.",
  "R1 minor: removed the repeated statement about adjacent transitions and pointed to the new clinical-significance stratification."),
 ("Contributions by new submitters (concordant new submissions) accounted for 127,652 (18.3%) of the total variant classification events",
  "New-submitter-driven reclassifications—where an added submitter shifted the aggregate category without producing a formal conflict—accounted for 127,652 (18.3%) of the total variant classification events",
  "R1 major: these events change the consensus category; 'concordant' was misleading."),
 ("HGVS_p in 78%, the latter reflecting its applicability only to protein-coding variants",
  "HGVS_p in 78% overall; because HGVS_p is defined only where a protein consequence is meaningful, completeness among the 2.81M coding variants is 99.7%, and the 78% reflects the inclusion of non-coding, intronic, and structural variants for which no protein consequence applies",
  "R2 minor: report HGVS_p completeness among variants where a protein consequence is meaningful (99.7%), not the raw 78% that mixes in non-coding variants."),
 ("Differences in aggregate usage between pathogenic and benign criteria at each strength level reflect differences in the types of evidence available for each classification direction rather than an asymmetry in evidentiary rigor.",
  "Pathogenic and benign criteria differ in aggregate usage because the two directions draw on different evidence types (e.g., PVS1 has no benign counterpart; BA1 has no pathogenic counterpart), not because one direction is held to a more rigorous evidentiary standard. Criteria usage also varied systematically by classification category, with pathogenic criteria concentrated in P/LP variants and benign criteria in B/LB variants (Figure 14d).",
  "R2 minor: sentence rewritten for clarity (and delete the stray quotation mark after 'rigor.'). Also adds an explicit in-text reference to Figure 14d with a distinct message, so 14d is no longer unreferenced (R2 minor)."),
 ("(e) Schema change taxonomy classifying each release as additive, deprecated, or mixed",
  "(e) Schema change taxonomy classifying each release as additive (new elements only), deprecated (removed elements only), or mixed (both added and removed elements), and as a rename when a new element name substring-matches a deprecated one",
  "R1 minor: define additive / deprecated / mixed in the Figure 4e legend."),
 ("Among the variants that underwent reclassification, the median time from first ClinVar appearance to first classification change was 2.58 years, with 24.5% of variants reclassified within one year and 81.8% reclassified within 5 years",
  "Because variants enter ClinVar at different times with unequal follow-up, we estimated the time to first reclassification with a right-censored Kaplan–Meier analysis over all 3,678,303 variants: the cumulative incidence of first reclassification was 2.5% (95% CI 2.53–2.56) at one year, 15.5% (15.47–15.59) at five years, and 30.9% at ten years, and the median was not reached because fewer than half of variants are ever reclassified. Among the subset that were reclassified, the median time to first change was 2.58 years",
  "R2 major #2 (editor-flagged): the time-to-event is now censored over ALL variants, not only reclassified ones. The prior 24.5%/81.8% figures were conditional on being reclassified and overstated the population tempo; the censored cumulative incidence is far lower. Add the new Kaplan–Meier figure."),
 ("The open-source analysis pipeline and processed checkpoint files are available upon request.",
  "The open-source analysis pipeline is available on GitHub (<URL>) and archived on Zenodo (<DOI>). Processed per-snapshot checkpoint data and analysis outputs are deposited on Zenodo (<DOI>). All source data are the publicly available NCBI ClinVar monthly XML releases (https://ftp.ncbi.nlm.nih.gov/pub/clinvar/xml/).",
  "Editor requirement: 'available upon request' does not meet Genome Medicine's data policy. Replace with GitHub + Zenodo DOIs (fill in the URLs before submission)."),
]

INSERTIONS = [
 ("attribution framework used in this study",
  ["Throughout, we use “reclassification” for any change in a variant’s aggregate consensus category between snapshots; “reinterpretation” for an evidence-driven change made by a submitter; and “administrative/aggregation change” for events (e.g., schema transition, a new concordant submitter, review-status relabeling) that alter metadata. Where a change crosses the P/LP–VUS–B/LB actionability boundary we describe it as medically significant, following ClinVar’s conflict-severity definitions."],
  "R1 minor: explicit terminology (reclassification vs reinterpretation vs administrative), applied consistently and addressing the 4.4/4.7 conflation."),
 ("rather than drastic reversals",
  ["To avoid conflating administrative and clinically consequential changes, we stratified the 342,057 reclassified variants (9.3%) by the most clinically significant change each underwent (Fig X). Every counted event involved a genuine category change (0 of 697,698 events were category-stable); concordant new submissions that leave the consensus unchanged are not included. Only 25,105 variants (0.68% of all variants; 7.3% of reclassified) ever underwent a medically significant change crossing the actionability boundary. A further 272,477 (7.41%) involved a composite or uncertain category (predominantly VUS↔Conflicting), 28,931 (0.79%) were non-medically-significant (VUS↔LB/B), and 15,544 (0.42%) involved only within-tier adjacent moves (P↔LP, B↔LB) that ClinVar considers concordant."],
  "R1 major & R2 major #1: new tiered-reclassification result (insert new Figure). Distinguishes clinically significant reclassification from any category change."),
 ("indicating greater research scrutiny",
  ["These differences were statistically significant with substantial effect sizes: cancer versus non-cancer genes, risk ratio (RR) 2.03 (95% CI 2.01–2.04); higher-prevalence versus rare Mendelian variants, RR 2.13 (2.08–2.18); and, by domain versus Other/Unclassified, RR 1.83–2.51 (all χ² p < 10⁻³⁰⁰). In a multivariable logistic regression adjusting for baseline submission count, baseline review status, follow-up time, submission era, variant class, and gene size, cancer-gene status remained associated with reclassification (adjusted OR 1.48, 95% CI 1.47–1.50) and all five disease domains remained significantly elevated (adjusted OR 1.20–1.55); domain differences are thus attenuated but not explained by testing intensity."],
  "R2 major #4: statistical tests with effect sizes + multivariable adjustment for the confounders the reviewer named (using BASELINE covariates to avoid a collider)."),
 ("the proportion of multi-submitter variants reaching consensus is improving",
  ["In segmented interrupted time-series regressions on the monthly snapshots, the VUS proportion showed significant level and slope changes at both the 2015 ACMG/AMP guideline release and the 2020 VCV transition (R²=0.90), and VCV-era inter-submitter concordance increased significantly at +2.9%/year (p=4×10⁻¹³). Because ClinVar’s “conflicting” review vocabulary is only interpretable from 2020 onward, concordance/conflict series are restricted to the VCV era; early-period VUS values partly reflect ClinVar’s initially sparse content, so the 2015 change is presented descriptively rather than as sole causal attribution."],
  "R2 major #5: interrupted time-series across the 2015 ACMG/AMP and 2020 VCV milestones, with honest scope caveats."),
 ("R01LM014344",
  ["Acknowledgements", "Not applicable."],
  "Editor: the Declarations were missing the 'Acknowledgements' subheading (required). Replace 'Not applicable' with any acknowledgements as appropriate."),
]

def main():
    doc = Document(DOCX)
    print("Applying replacements:")
    for find, repl, com in REPLACEMENTS:
        tracked_replace(doc, find, repl, com)
    print("Applying insertions:")
    for anchor, paras, com in INSERTIONS:
        insert_after(doc, anchor, paras, com)
    doc.save(DOCX)
    print("\nSaved. Verifying...")
    d2 = Document(DOCX)
    import zipfile
    z = zipfile.ZipFile(DOCX)
    xml = z.read('word/document.xml').decode('utf-8', 'ignore')
    print(f"  w:ins occurrences: {xml.count('<w:ins ')}")
    print(f"  w:del occurrences: {xml.count('<w:del ')}")
    print(f"  comments part present: {'word/comments.xml' in z.namelist()}")
    if 'word/comments.xml' in z.namelist():
        c = z.read('word/comments.xml').decode('utf-8', 'ignore')
        print(f"  comments: {c.count('<w:comment ')}")
    # confirm key new phrases present
    full = "\n".join(p.text for p in d2.paragraphs)
    for probe in ["only 0.68% (25,105 variants)", "adjusted OR 1.48", "+2.9%/year", "99.7%", "new-submitter-driven"]:
        print(f"  contains {probe[:32]!r}: {probe in full}")

if __name__ == "__main__":
    main()
